"""Load versioned channel artifacts and execute MLP prediction."""

from __future__ import annotations

import json
import threading
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import joblib
import numpy as np
import pandas as pd
import torch
from sklearn.preprocessing import StandardScaler

from config import (DATA_DIR, TRAIN_DATA_DIR, PREDICTION_ALLOWED_DIRS,
                    MAX_PREDICTION_FILE_BYTES, MAX_BATCH_PREDICTIONS, APP_ENV)
from services.resource_gate import ResourceGate

from models.mlp import DynamicMLP, MLPConfig


class InputFileChangedError(ValueError):
    pass


@dataclass
class LoadedArtifact:
    model: DynamicMLP
    scaler_x: StandardScaler
    scaler_y: StandardScaler
    config: MLPConfig
    metrics: dict
    version_dir: Path | None = None


class PredictionService:
    def __init__(
        self,
        model_dir: str | Path,
        device: str | torch.device | None = None,
    ) -> None:
        self.model_dir = Path(model_dir)
        self.device = torch.device(device or ("cuda" if torch.cuda.is_available() else "cpu"))
        self._cache: dict[int, LoadedArtifact] = {}
        self._lock = threading.RLock()
        self.resource_gate: ResourceGate | None = None
        configured = [Path(item).expanduser().resolve() for item in PREDICTION_ALLOWED_DIRS.split(";") if item.strip()]
        self.allowed_dirs = configured if APP_ENV=='production' else configured or [self.model_dir.resolve().parent, DATA_DIR, TRAIN_DATA_DIR]

    def checked_path(self, file_path: str) -> Path:
        path = Path(file_path).expanduser().resolve()
        if not any(path.is_relative_to(root) for root in self.allowed_dirs):
            raise ValueError("预测文件不在允许的数据目录内")
        if path.suffix.lower() not in {".csv", ".xlsx"}:
            raise ValueError("预测文件只支持 .csv 和 .xlsx")
        if not path.is_file():
            raise FileNotFoundError(f"预测文件不存在: {path}")
        if path.stat().st_size > MAX_PREDICTION_FILE_BYTES:
            raise ValueError("预测文件超过大小上限")
        return path

    def batch_work(self, files: list[dict]) -> int:
        total = 0
        for item in files:
            candidate = Path(item["file_path"]).expanduser().resolve()
            if not candidate.is_file() or candidate.suffix.lower() not in {".csv", ".xlsx"}:
                continue  # The task records these ordinary input errors with file/channel context.
            path = self.checked_path(item["file_path"])
            if path.suffix.lower() == ".csv":
                with path.open("rb") as stream:
                    rows = max(sum(1 for _ in stream) - 1, 0)
            else:
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    sheet = item.get("sheet_name", 0)
                    worksheet = workbook.worksheets[sheet] if isinstance(sheet, int) else workbook[sheet]
                    rows = max(worksheet.max_row - 1, 0)
                finally:
                    workbook.close()
            if not 1 <= rows <= 100_000:
                raise ValueError("预测文件必须包含 1 到 100000 行")
            total += rows * len(item["channels"])
            if total > MAX_BATCH_PREDICTIONS:
                raise ValueError("批量预测总条数超过上限")
        return total

    def artifact_dir(self, channel: int) -> Path:
        return self.model_dir / f"channel_{channel}"

    def active_artifact_dir(self, channel: int) -> Path:
        root = self.artifact_dir(channel)
        pointer = root / "CURRENT"
        if not pointer.is_file():
            return root  # Existing deployed model layout.
        version = pointer.read_text(encoding="utf-8").strip()
        if not version or Path(version).name != version:
            raise ValueError("模型版本指针无效")
        return root / "versions" / version

    def available_channels(self) -> list[int]:
        channels: list[int] = []
        for channel in range(8):
            if self._has_complete_artifact(channel):
                channels.append(channel)
        return channels

    def invalidate(self, channel: int) -> None:
        with self._lock:
            self._cache.pop(channel, None)

    def predict(self, channel: int, rows: list[list[float]] | np.ndarray,
                event_callback: Callable[[str, str], None] | None = None) -> list[dict[str, float]]:
        emit = event_callback or (lambda operation, message: None)
        emit("preprocess", "校验并转换四维输入")
        values = np.asarray(rows, dtype=np.float32)
        if values.ndim != 2 or values.shape[0] == 0 or values.shape[1] != 4:
            raise ValueError("data 必须是非空的 N x 4 数值数组")
        if not np.isfinite(values).all():
            raise ValueError("data 不能包含 NaN 或无穷大")
        emit("load_model", f"加载 CH{channel} 模型")
        artifact = self.load(channel)
        emit("inference", f"开始 CH{channel} 模型推理")
        transformed = artifact.scaler_x.transform(values).astype(np.float32)
        tensor = torch.as_tensor(transformed, dtype=torch.float32, device=self.device)
        if self.resource_gate is None:
            with torch.inference_mode():
                scaled_predictions = artifact.model(tensor).cpu().numpy()
        else:
            with self.resource_gate.enter(), torch.inference_mode():
                scaled_predictions = artifact.model(tensor).cpu().numpy()
        emit("postprocess", f"CH{channel} 预测后处理")
        predictions = artifact.scaler_y.inverse_transform(scaled_predictions)
        return [
            {"ea": float(prediction[0]), "bias": float(prediction[1])}
            for prediction in predictions
        ]

    def predict_file(
        self,
        channel: int,
        file_path: str,
        sheet_name: str | int = 0,
        feature_columns: list[str | int] | None = None,
        event_callback: Callable[[str, str], None] | None = None,
        batch_sink: Callable[[list[dict[str, float]]], None] | None = None,
        collect: bool = True,
        input_size_bytes: int | None = None,
        input_mtime_ns: int | None = None,
    ) -> tuple[list[dict[str, float]], dict]:
        emit = event_callback or (lambda operation, message: None)
        emit("load_file", f"读取预测文件 {Path(file_path).name}")
        path = self.checked_path(file_path)
        def check_snapshot():
            stat = path.stat()
            if (input_size_bytes is not None and stat.st_size != input_size_bytes) or (
                    input_mtime_ns is not None and stat.st_mtime_ns != input_mtime_ns):
                raise InputFileChangedError("文件在识别后发生变化，请刷新文件夹后重新预测")
        check_snapshot()
        if path.suffix.lower() == ".csv":
            frame = None
            for encoding in ("utf-8-sig", "gb18030", "latin1"):
                try:
                    frame = pd.read_csv(path, encoding=encoding, nrows=100_001)
                    break
                except UnicodeDecodeError:
                    continue
            if frame is None:
                raise ValueError("无法识别 CSV 编码")
        elif path.suffix.lower() == ".xlsx":
            frame = pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl")
        else:
            raise ValueError("预测文件只支持 .csv 和 .xlsx")
        check_snapshot()
        emit("validate", f"校验文件 {path.name}")
        if frame.empty or len(frame) > 100_000:
            raise ValueError("预测文件必须包含 1 到 100000 行")
        if feature_columns is None:
            if frame.shape[1] == 4:
                selected = frame.iloc[:, :4]
            elif frame.shape[1] >= 7:
                selected = frame.iloc[:, [3, 4, 5, 6]]
            else:
                raise ValueError("文件需要恰好 4 列或至少 7 列")
        elif all(isinstance(column, str) for column in feature_columns):
            selected = frame.loc[:, feature_columns]
        elif all(isinstance(column, int) and not isinstance(column, bool) for column in feature_columns):
            selected = frame.iloc[:, feature_columns]
        else:
            raise ValueError("feature_columns 必须全部是列名或全部是列序号")
        emit("preprocess", f"提取 {path.name} 的四维特征")
        numeric = selected.apply(pd.to_numeric, errors="coerce")
        if numeric.isna().any().any():
            raise ValueError("预测特征包含空值或非数值")
        from config import PREDICTION_CHUNK_ROWS
        values = numeric.to_numpy(dtype=np.float32)
        with_rows = []
        for start in range(0, len(values), max(1, PREDICTION_CHUNK_ROWS)):
            predictions = self.predict(channel, values[start:start + max(1, PREDICTION_CHUNK_ROWS)],
                                       event_callback=event_callback if start == 0 else None)
            batch = [{"row": start + index, **prediction}
                     for index, prediction in enumerate(predictions, start=2)]
            if batch_sink:
                batch_sink(batch)
            if collect:
                with_rows.extend(batch)
        return with_rows, {
            "file_name": path.name,
            "feature_columns": [str(column) for column in selected.columns],
            "row_count": len(frame),
        }

    def load(self, channel: int) -> LoadedArtifact:
        if isinstance(channel, bool) or not 0 <= channel <= 7:
            raise ValueError("channel 必须是 0 到 7 的整数")
        with self._lock:
            directory = self.active_artifact_dir(channel)
            files = self._artifact_files(directory)
            missing = [path.name for path in files if not path.is_file()]
            if missing:
                raise FileNotFoundError(
                    f"Channel {channel} 模型 artifact 不完整，缺少: {', '.join(missing)}"
                )
            cached = self._cache.get(channel)
            if cached is not None and cached.version_dir == directory:
                return cached
            loaded = self._load(directory)
            loaded.version_dir = directory
            self._cache[channel] = loaded
            return loaded

    def _has_complete_artifact(self, channel: int) -> bool:
        try:
            return all(path.is_file() for path in self._artifact_files(self.active_artifact_dir(channel)))
        except (OSError, ValueError):
            return False

    @staticmethod
    def _artifact_files(directory: Path) -> tuple[Path, ...]:
        return (
            directory / "model.pth",
            directory / "config.json",
            directory / "scaler_x.joblib",
            directory / "scaler_y.joblib",
            directory / "metrics.json",
        )

    def _load(self, directory: Path) -> LoadedArtifact:
        config_values = json.loads((directory / "config.json").read_text(encoding="utf-8"))
        config = MLPConfig.from_dict(config_values)
        if config.input_dim != 4 or config.output_dim != 2:
            raise ValueError("当前 API 只支持 4 维输入、2 维输出")
        model = DynamicMLP(config)
        try:
            state_dict = torch.load(directory / "model.pth", map_location="cpu", weights_only=True)
        except TypeError:  # PyTorch 1.x compatibility
            state_dict = torch.load(directory / "model.pth", map_location="cpu")
        model.load_state_dict(state_dict)
        model.to(self.device).eval()
        scaler_x = joblib.load(directory / "scaler_x.joblib")
        scaler_y = joblib.load(directory / "scaler_y.joblib")
        if getattr(scaler_x, "n_features_in_", None) != 4:
            raise ValueError("scaler_x 与 4 维输入不兼容")
        if getattr(scaler_y, "n_features_in_", None) != 2:
            raise ValueError("scaler_y 与 2 维输出不兼容")
        metrics = json.loads((directory / "metrics.json").read_text(encoding="utf-8"))
        return LoadedArtifact(model, scaler_x, scaler_y, config, metrics)
